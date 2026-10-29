"""
Datos del visor Relieve3D.html a partir del paquete de capas de GeoIngenIA
(../Capas_Rio_Sinu_GeoIngenIA/, 234 MB, fuera de git).

Escribe en ../datos/:
  dem.png, direcciones.png, zonas.png, rejilla.json, subcuencas.geojson, drenaje.geojson
      lo que usa el motor 3D (../../motor3d/), con el DEM ALOS PALSAR, las 57 subcuencas
      del POMCA + el resto de la cuenca (valle y cauce principal) y la red de drenaje oficial
  capas/*.geojson      cuenca, estaciones IDEAM, infraestructura, lotico/lentico 2020,
                       areas de interes ambiental y desastres (todo en lat/lon)
  agua/inundacion_AAAA.png, agua/sequia_AAAA.png   las 82 imagenes pasadas a lat/lon
  indices/ndvi_AAAA.png (y ndwi, mndwi)   rasteres anuales de Indices_Espectrales_Geovisor a 250 m
  series.json          caudal y lluvia de las estaciones, area de agua por ano, ONI/RONI,
                       lotico/lentico por subcuenca

Uso (desde la raiz del repo):
  py -m pip install pyshp numpy pillow tifffile imagecodecs pyproj shapely openpyxl
  py "Cuencas/Rio Sinu/herramientas/preparar_geoingenia.py"
"""
import csv
import glob
import json
import os
import sys

import numpy as np
import openpyxl
import shapefile
import shapely
import tifffile
from PIL import Image
from pyproj import CRS, Transformer
from shapely.geometry import mapping, shape
from shapely.ops import unary_union

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.normpath(os.path.join(AQUI, "..", "..", ".."))
PAQ = os.path.join(AQUI, "..", "Capas_Rio_Sinu_GeoIngenIA")
SALIDA = os.path.join(AQUI, "..", "datos")
sys.path.insert(0, os.path.join(RAIZ, "Cuencas", "Cundinamarca", "herramientas"))
import preparar_datos as motor  # noqa: E402

FACTOR = 7              # DEM ALOS de 12,5 m -> ~87 m
ORDEN_MIN = 3           # orden de Strahler minimo de la red oficial que se dibuja y se graba
MARGEN_PNG = 20000      # las PNG de agua cubren la cuenca + 20 km (ver georreferenciar_png)
GRUPOS = {"1302": "Medio Sinú", "1303": "Bajo Sinú"}
GRUPO_VALLE = "Alto Sinú y cauce principal"


# ---------------------------------------------------------------------
# Lectura de shapefiles (cada uno viene en un sistema distinto)
# ---------------------------------------------------------------------
def arreglar(t):
    """Los .dbf vienen unos en UTF-8 y otros en Latin-1: se leen como Latin-1 y si el
    texto era UTF-8 se recupera."""
    if not isinstance(t, str):
        return t
    try:
        t = t.encode("latin-1").decode("utf-8")
    except (UnicodeDecodeError, UnicodeEncodeError):
        pass
    return t.strip()


def leer(rel, simplificar=0.0, recorte=None):
    """Lista de (propiedades, geometria en lat/lon). simplificar en grados."""
    shp = os.path.join(PAQ, rel)
    sf = shapefile.Reader(shp, encoding="latin-1")
    crs = CRS.from_wkt(open(shp[:-4] + ".prj", encoding="latin-1").read())
    t = Transformer.from_crs(crs, 4326, always_xy=True)
    campos = [arreglar(f[0]) for f in sf.fields[1:]]
    salida = []
    for sr in sf.shapeRecords():
        if sr.shape.shapeType == 0 or not sr.shape.points:
            continue
        g = shapely.transform(shape(sr.shape.__geo_interface__),
                              lambda xy: np.column_stack(t.transform(xy[:, 0], xy[:, 1])))
        if g.geom_type.endswith("Polygon"):
            g = shapely.make_valid(g)
            g = shapely.union_all([p for p in shapely.get_parts(g) if p.geom_type.endswith("Polygon")])
        if recorte is not None:
            if not g.intersects(recorte):
                continue
            if g.geom_type.endswith("Polygon"):
                g = g.intersection(recorte)
        if simplificar:
            g = g.simplify(simplificar, preserve_topology=True)
        if g.is_empty:
            continue
        g = shapely.transform(g, lambda xy: np.round(xy, 5))
        salida.append(({c: arreglar(v) for c, v in zip(campos, sr.record)}, g))
    return salida


def num(v, d=2):
    try:
        return round(float(v), d)
    except (TypeError, ValueError):
        return None


def guardar(nombre, feats):
    os.makedirs(os.path.join(SALIDA, "capas"), exist_ok=True)
    fc = {"type": "FeatureCollection", "features": [
        {"type": "Feature", "properties": {k: v for k, v in p.items() if v not in (None, "", "None")},
         "geometry": mapping(g)} for p, g in feats]}
    texto = json.dumps(fc, ensure_ascii=False, separators=(",", ":"))
    open(os.path.join(SALIDA, "capas", nombre + ".geojson"), "w", encoding="utf-8").write(texto)
    print(f"  capas/{nombre}: {len(feats)} elementos, {len(texto) / 1e6:.2f} MB")


# ---------------------------------------------------------------------
# Motor 3D: DEM, subcuencas y drenaje
# ---------------------------------------------------------------------
def leer_dem():
    with tifffile.TiffFile(os.path.join(PAQ, "01_Cuenca_Drenaje_DEM", "DEM", "DEM_ALOS_MACROCUENCA.tif")) as t:
        tags = t.pages[0].tags
        x0, y0 = tags["ModelTiepointTag"].value[3:5]
        paso = tags["ModelPixelScaleTag"].value[0]
        a = t.pages[0].asarray()
    h, w = (a.shape[0] // FACTOR) * FACTOR, (a.shape[1] // FACTOR) * FACTOR
    dem = a[:h, :w].reshape(h // FACTOR, FACTOR, w // FACTOR, FACTOR).mean(axis=(1, 3)).astype(np.float32)
    dem[dem < -100] = np.nan
    meta = {"oeste": x0, "norte": y0, "dx": paso * FACTOR, "dy": paso * FACTOR, "ancho": dem.shape[1], "alto": dem.shape[0]}
    meta["este"] = meta["oeste"] + meta["ancho"] * meta["dx"]
    meta["sur"] = meta["norte"] - meta["alto"] * meta["dy"]
    return dem, meta


def motor3d(cuenca, info_agua):
    ini = motor.t0()
    dem, meta = leer_dem()
    motor.log(f"DEM ALOS {meta['ancho']}x{meta['alto']} a {meta['dx'] * 111320:.0f} m", ini)
    a_9377 = Transformer.from_crs(4326, 9377, always_xy=True)
    km = lambda g: shapely.transform(g, lambda xy: np.column_stack(a_9377.transform(xy[:, 0], xy[:, 1])))

    subs = leer("01_Cuenca_Drenaje_DEM/Subcuencas/subcuencas.shp")
    # el resto de la cuenca (valle y cauce principal) no tiene poligono en el POMCA
    resto = cuenca.difference(unary_union([g.buffer(0.0005) for _, g in subs])).buffer(-0.0004).buffer(0)
    resto = shapely.union_all([p for p in shapely.get_parts(resto) if km(p).area > 5e6])
    subs.append(({"NOM_NSS": "Valle del Sinú y cauce principal", "SZH": "", "COD_POMCA": ""}, resto))

    zonas = []
    for k, (p, g) in enumerate(subs, start=1):
        partes = list(getattr(g, "geoms", [g]))
        gk = km(g)
        extra = {"pomca": p.get("POMCA"), "t_corrien": p.get("T_CORRIEN"), "caud_med": num(p.get("CAUD_MED")),
                 "q_max": num(p.get("Q_MAX")), "q_min": num(p.get("Q_MIN")), "q_ambienta": num(p.get("Q_AMBIENTA")),
                 "q_disponib": num(p.get("Q_DISPONIB")), "rend_hid": num(p.get("REND_HID")),
                 "tiem_conc": num(p.get("TIEM_CONC")), "den_dren": num(p.get("DEN_DREN"))}
        if k == len(subs):
            extra = {"nota": "Resto de la cuenca: el POMCA no le da un polígono propio."}
        zonas.append({
            "id": k, "codch": int(p.get("SZH") or 1301), "cuenca": GRUPOS.get(str(p.get("SZH")), GRUPO_VALLE),
            "codsch": p.get("POMCA") or "—", "nombre": p.get("NOM_NSS") or p.get("COD_POMCA"),
            "anillos": [list(pp.exterior.coords) for pp in partes], "huecos": sum(len(pp.interiors) for pp in partes),
            "geometria": mapping(g.simplify(0.00015, preserve_topology=True)),
            "area_km2": round(gk.area / 1e6, 1), "perimetro_km": round(gk.length / 1000, 1),
            "extra": {a: b for a, b in extra.items() if b not in (None, "")},
        })
    zonas.sort(key=lambda z: -z["huecos"])

    red = red_sin_rectas(info_agua)
    lineas = []
    for _, g in red:
        lineas += [list(l.coords) for l in getattr(g, "geoms", [g])]
    motor.log(f"{len(zonas)} subcuencas, red oficial: {len(red)} tramos de orden >= {ORDEN_MIN}", ini)
    motor.procesar(dem, meta, zonas, lineas, SALIDA, ini, {
        "fuente": "ALOS PALSAR 12,5 m (GeoIngenIA), remuestreado a ~87 m",
        "rios_grabados": f"Red de drenaje oficial GeoIngenIA (orden >= {ORDEN_MIN})",
        "variables": ("DATA_SUBCUENCAS_SINU", "DATA_DRENAJE_SINU")})
    # la red que se dibuja es la oficial (no la calculada), con su orden de Strahler real
    feats = []
    for p, g in red:
        g = g.simplify(0.0003, preserve_topology=True)
        for l in getattr(g, "geoms", [g]):
            feats.append({"type": "Feature", "properties": {"orden": int(p["grid_code"])},
                          "geometry": {"type": "LineString", "coordinates": [[round(x, 5), round(y, 5)] for x, y in l.coords]}})
    feats.sort(key=lambda f: f["properties"]["orden"])
    texto = json.dumps({"type": "FeatureCollection", "features": feats}, separators=(",", ":"))
    open(os.path.join(SALIDA, "drenaje.geojson"), "w", encoding="utf-8").write(texto)
    for viejo in ("drenaje.js", "subcuencas.js"):
        if os.path.exists(os.path.join(SALIDA, viejo)):
            os.remove(os.path.join(SALIDA, viejo))
    motor.log(f"drenaje oficial: {len(feats)} tramos, {len(texto) / 1e6:.2f} MB", ini)


def mascara_cienagas():
    """Cuerpos de agua anchos que se inundan seguido (cienagas, Urra): agua en al menos 4 de
    los 41 anos, sin los cauces delgados (apertura de 3 pixeles) y con 7 pixeles (~800 m) de margen."""
    acc = None
    for f in glob.glob(os.path.join(SALIDA, "agua", "inundacion_*.png")):
        a = np.array(Image.open(f))[..., 3] > 0
        acc = a.astype(np.int16) if acc is None else acc + a

    def mover(a, df, dc):
        b = np.zeros_like(a)
        H, W = a.shape
        b[max(df, 0):H + min(df, 0), max(dc, 0):W + min(dc, 0)] = a[max(-df, 0):H + min(-df, 0), max(-dc, 0):W + min(-dc, 0)]
        return b

    def morf(a, r, erosion):
        o = a.copy()
        for df in range(-r, r + 1):
            for dc in range(-r, r + 1):
                if df * df + dc * dc <= r * r:
                    o = (o & mover(a, df, dc)) if erosion else (o | mover(a, df, dc))
        return o
    return morf(morf(acc >= 4, 3, True), 7, False)


def red_sin_rectas(info_agua):
    """Red oficial (orden >= ORDEN_MIN) sin las lineas falsas de las zonas planas.

    La red viene de un calculo sobre el DEM; en las cienagas del bajo Sinu ese calculo
    manda el agua en lineas rectas paralelas que no existen. Se quitan:
      - los tramos de mas de 1,5 km casi perfectamente rectos (aun alisando los escalones);
      - lo que cae dentro de los cuerpos de agua anchos (mascara_cienagas), donde no hay cauce."""
    import math
    from shapely.geometry import LineString
    caja = info_agua["caja"]
    agua_ancha = mascara_cienagas()
    H, W = agua_ancha.shape
    px = (caja[2] - caja[0]) / W

    def en_agua(x, y):
        c, f = int((x - caja[0]) / px), int((caja[3] - y) / px)
        return 0 <= c < W and 0 <= f < H and agua_ancha[f, c]

    shp = os.path.join(PAQ, "01_Cuenca_Drenaje_DEM", "Drenaje_orden8", "Drenaje_Orden8_9377.shp")
    sf = shapefile.Reader(shp)
    t = Transformer.from_crs(9377, 4326, always_xy=True)
    i_orden = [f[0] for f in sf.fields[1:]].index("grid_code")
    red, rectas, en_cienaga = [], 0.0, 0
    for sh, r in zip(sf.shapes(), sf.records()):
        orden = int(r[i_orden])
        if orden < ORDEN_MIN:
            continue
        pts = sh.points
        linea = LineString(pts)
        if linea.length >= 1500 and math.dist(pts[0], pts[-1]) / linea.simplify(60).length > 0.993:
            rectas += linea.length
            continue
        lon, lat = t.transform([p[0] for p in pts], [p[1] for p in pts])
        ll = list(zip(lon, lat))
        dentro = [en_agua(x, y) for x, y in ll]
        pedazo = [ll[0]]
        for k in range(1, len(ll)):
            if dentro[k - 1] and dentro[k]:
                en_cienaga += 1
                if len(pedazo) > 1:
                    red.append(({"grid_code": orden}, LineString(pedazo)))
                pedazo = [ll[k]]
            else:
                pedazo.append(ll[k])
        if len(pedazo) > 1:
            red.append(({"grid_code": orden}, LineString(pedazo)))
    print(f"  red oficial: {rectas / 1000:.0f} km de rectas falsas y {en_cienaga} segmentos dentro de cienagas quitados")
    return red


# ---------------------------------------------------------------------
# Capas tematicas
# ---------------------------------------------------------------------
def capas(cuenca, area_km2):
    caja = cuenca.envelope.buffer(0.15)
    guardar("cuenca", [({"nombre": "Cuenca del río Sinú", "area_km2": area_km2}, cuenca.simplify(0.0001))])

    # Estaciones IDEAM (Web Mercator); el codigo se normaliza sin ceros a la izquierda
    est, vistos = [], set()
    for p, g in leer("02_Estaciones_IDEAM/Estaciones_Cuenca_IDEAM.shp"):
        if str(p["Codigo"]).lstrip("0") in vistos:  # el shapefile trae cada estacion dos veces
            continue
        vistos.add(str(p["Codigo"]).lstrip("0"))
        est.append(({"codigo": str(p["Codigo"]).lstrip("0"), "nombre": p["Nombre"], "categoria": p["Categoria"],
                     "tecnologia": p["Tecnologia"], "estado": p["Estado"], "municipio": p["Municipio"],
                     "departamento": p["Departamen"], "altitud": num(p["Altitud"], 0), "corriente": p.get("Corriente"),
                     "instalada": p.get("Fecha_inst"), "suspendida": p.get("Fecha_susp"), "entidad": p.get("Entidad")}, g))

    # series de las estaciones
    def tabla(rel):
        with open(os.path.join(PAQ, "02_Estaciones_IDEAM", rel), encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))
    series = {}
    for rel, clave, var, mensual in [("caudal_m3s_por_estacion_mensual.csv", "qm", "caudal_m3s_promedio", True),
                                     ("caudal_m3s_por_estacion_anual.csv", "qa", "caudal_m3s_promedio", False),
                                     ("precipitacion_mm_por_estacion_mensual.csv", "pm", "precipitacion_mm_promedio", True),
                                     ("precipitacion_mm_por_estacion_anual.csv", "pa", "precipitacion_mm_promedio", False)]:
        for r in tabla(rel):
            c = r["CodigoEstacion"].lstrip("0")
            fila = [int(r["year"])] + ([int(r["month"])] if mensual else []) + [num(r[var], 2)]
            series.setdefault(c, {}).setdefault(clave, []).append(fila)
    for p, _ in est:
        s = series.get(p["codigo"], {})
        p["con_caudal"] = "qm" in s or None
        p["con_lluvia"] = "pm" in s or None
    guardar("estaciones", est)

    # infraestructura
    guardar("embalse", [({"nombre": p["NOMBRE_GEO"]}, g) for p, g in leer("06_Infraestructura/Embalses cuenca rio sinu.shp", 0.0001)])
    guardar("lagunas", [({"nombre": p.get("NOMBRE_GEO"), "area_ha": num(float(p["SHAPE_Area"]) / 1e4, 1)}, g)
                        for p, g in leer("06_Infraestructura/Lagunas Cuenca Rio sinu.shp", 0.00005)])
    guardar("plantas", [({"nombre": p["PTAP"], "empresa": p["EMPRESA"], "municipio": p["MpNombre"]}, g)
                        for p, g in leer("06_Infraestructura/Plantas tratamiento cuenca rio sinu.shp")])
    guardar("poblaciones", [({"nombre": p["zu_cnmbre"], "altitud": num(p["zu_naltd"], 0)}, g)
                            for p, g in leer("06_Infraestructura/Poblaciones cuenca rio sinu.shp", 0.00005)])
    guardar("puentes", [({"nombre": p["nombre"], "via": p["carretera"], "material": p.get("material"),
                          "longitud_m": num(p.get("luztotal"), 1)}, g) for p, g in leer("06_Infraestructura/Puentes cuenca rio sinu.shp")])
    guardar("rellenos", [({"nombre": p.get("NOMBRE_SIT"), "municipio": p.get("MUNICIPIO"), "tipo": p.get("TIPO"),
                           "toneladas_dia": num(p.get("TONELADA_D"), 1)}, g)
                         for p, g in leer("06_Infraestructura/Relleno sanitario Cuenca Rio sinu.shp", 0.0001)])
    guardar("vias", [({"nombre": p["nombreruta"], "tramo": p.get("sector")}, g)
                     for p, g in leer("06_Infraestructura/Vias cuenca rio sinu.shp", 0.0002)])

    # desastres (hoja de calculo georreferenciada)
    from shapely.geometry import Point
    wb = openpyxl.load_workbook(os.path.join(PAQ, "06_Infraestructura", "desastres_naturales_cordoba_1980_2026_georreferenciado.xlsx"), read_only=True)
    filas = list(wb.worksheets[0].iter_rows(values_only=True))
    des = []
    for r in filas[1:]:
        if r[6] is None or r[7] is None:
            continue
        des.append(({"titulo": r[0], "anio": r[1], "mes": r[2], "tipo": r[3], "lugar": r[5], "fuente": r[8], "url": r[9]},
                    Point(round(float(r[7]), 5), round(float(r[6]), 5))))
    guardar("desastres", des)

    # lotico / lentico 2020
    guardar("lotico_lentico_2020", [({"tipo": p["tipo"], "permanencia": p["perm"], "area_ha": num(p["area_ha"], 2)}, g)
                                    for p, g in leer("07_Lotico_Lentico/2020/sinu_lotico_lentico_2020.shp", 0.00005)])

    # areas de interes ambiental
    A = "09_Areas_interes_ambiental/"
    guardar("clase_suelo", [({"clase": int(float(p["CLASESUE"])) if p["CLASESUE"] not in ("", None) else None, "area_ha": num(p["AREA_ha"], 1)}, g)
                            for p, g in leer(A + "Clase suelo/Clase de suelo.shp", 0.0011, caja)])
    guardar("dmi", [({"nombre": p["NOMBRE"], "acto": p.get("NUM_RES_VI"), "area_ha": num(p["AREA_RES"], 0)}, g)
                    for p, g in leer(A + "Distrito manejo integrado/Distrito de Manejo Integrado.shp", 0.0002)])
    guardar("parque_paramillo", [({"nombre": p["NOMBRE"], "resolucion": p.get("PN_RESOL"), "area_ha": num(p["AREA_RES"], 0)}, g)
                                 for p, g in leer(A + "Parque Natural/SINAP Parque Naciona.shp", 0.0002)])
    guardar("sinap", [({"nombre": p["NOM"], "categoria": p["PN_CAT"], "acto": p.get("NUM_RES_VI")}, g)
                      for p, g in leer(A + "Parque Natural/SINAP_PNN.shp", 0.0005, caja)])
    guardar("reserva_ley2", [({"nombre": "Zona de Reserva Forestal del Pacífico (Ley 2ª de 1959)", "area_ha": num(p["AREA_RES"], 0)}, g)
                             for p, g in leer(A + "Reserva forestal Ley 2/AECC_Reserva forestal Ley 2.shp", 0.0003, caja)])
    guardar("rnsc", [({"nombre": p["NOM_RNSC"], "acto": p.get("AC_ADT"), "area_ha": num(p["AREA_ha"], 1)}, g)
                     for p, g in leer(A + "Reserva Natural Sociedad Civil/Reserva Natural Sociedad Civil.shp", 0.0001)])
    guardar("conservacion_suelo", [({"nombre": p["NOMBRE"], "area_ha": num(p["AREA_RES"], 0)}, g)
                                   for p, g in leer(A + "Reserva suelo/Dsitrito de conservacion de suelo.shp", 0.0001)])
    return series


# ---------------------------------------------------------------------
# Inundacion y sequia por ano
# ---------------------------------------------------------------------
def georreferenciar_png(cuenca_9377):
    """Las PNG (1400 x 2291, EPSG:9377) no traen georreferencia. Calzando su agua con los
    poligonos de agua del paquete (Urra, lagunas, lotico/lentico 2020) se encontro que tres
    bordes coinciden a menos de 35 m con la caja de la cuenca + 20 km, con pixel cuadrado:
    oeste, este y norte salen de esa caja y el sur del tamano del pixel."""
    x0, _, x1, y1 = cuenca_9377.bounds
    oeste, este, norte = x0 - MARGEN_PNG, x1 + MARGEN_PNG, y1 + MARGEN_PNG
    s = (este - oeste) / 1400
    return oeste, norte, s


def agua(cuenca_9377):
    oeste, norte, s = georreferenciar_png(cuenca_9377)
    alto_png = 2291
    a84 = Transformer.from_crs(9377, 4326, always_xy=True)
    a9377 = Transformer.from_crs(4326, 9377, always_xy=True)
    xs = [oeste, oeste + 1400 * s]
    ys = [norte, norte - alto_png * s]
    lon, lat = a84.transform([xs[0], xs[1], xs[0], xs[1]], [ys[0], ys[0], ys[1], ys[1]])
    caja = [min(lon), min(lat), max(lon), max(lat)]
    paso = 0.00105  # ~117 m, como el original
    W, H = int((caja[2] - caja[0]) / paso), int((caja[3] - caja[1]) / paso)
    LO, LA = np.meshgrid(caja[0] + (np.arange(W) + 0.5) * paso, caja[3] - (np.arange(H) + 0.5) * paso)
    X, Y = a9377.transform(LO, LA)
    ci = np.floor((X - oeste) / s).astype(int)
    fi = np.floor((norte - Y) / s).astype(int)
    ok = (ci >= 0) & (ci < 1400) & (fi >= 0) & (fi < alto_png)
    ci, fi = np.where(ok, ci, 0), np.where(ok, fi, 0)
    os.makedirs(os.path.join(SALIDA, "agua"), exist_ok=True)
    anios = {}
    for tipo, carpeta in (("inundacion", "03_Inundacion_por_anio"), ("sequia", "04_Sequia_por_anio")):
        for f in sorted(glob.glob(os.path.join(PAQ, carpeta, "*", "*.png"))):
            anio = int(os.path.basename(os.path.dirname(f)))
            a = np.array(Image.open(f).convert("RGBA"))
            sal = a[fi, ci]
            sal[~ok] = 0
            Image.fromarray(sal, "RGBA").save(os.path.join(SALIDA, "agua", f"{tipo}_{anio}.png"), optimize=True)
            anios.setdefault(tipo, []).append(anio)
    tam = sum(os.path.getsize(p) for p in glob.glob(os.path.join(SALIDA, "agua", "*.png")))
    print(f"  agua: {sum(len(v) for v in anios.values())} imagenes {W}x{H}, {tam / 1e6:.1f} MB")
    return {"caja": [round(v, 5) for v in caja], "anios": anios}


# ---------------------------------------------------------------------
# Cobertura del suelo: modelo propio del proyecto (Random Forest sobre Landsat, 1985-2026)
# ---------------------------------------------------------------------
SUAVE = 3   # los mapas vienen a ~340 m: se agrandan 3 veces y se redondean los bordes por mayoria


def mayoria(a, nodata, f, r=2):
    """Agranda f veces y aplica un filtro de moda (ventana (2r+1)^2) que redondea las
    escaleras sin mezclar colores. Si un pixel del original desaparece del todo (un pueblo,
    una cienaga chica) se devuelve tal cual."""
    g = np.kron(a, np.ones((f, f), np.uint8))
    clases = [c for c in np.unique(a) if c != nodata]
    p = np.pad(g, r, mode="edge")
    mejor, cuenta = g.copy(), np.zeros(g.shape, np.int16)
    k = 2 * r + 1
    for c in clases:
        s = np.pad(np.cumsum(np.cumsum(p == c, 0), 1), ((1, 0), (1, 0)))
        n = s[k:, k:] - s[:-k, k:] - s[k:, :-k] + s[:-k, :-k]
        mas = n > cuenta
        mejor[mas], cuenta[mas] = c, n[mas]
    mejor[g == nodata] = nodata   # el borde de la cuenca no se mueve
    H, W = a.shape
    perdido = ~(mejor.reshape(H, f, W, f) == a[:, None, :, None]).any(axis=(1, 3))
    fi, ci = np.nonzero(perdido)
    for di in range(f):
        for dj in range(f):
            mejor[fi * f + di, ci * f + dj] = a[fi, ci]
    return mejor


def suelo():
    base = os.path.join(PAQ, "10_Cambio_uso_suelo")
    with open(os.path.join(base, "areas_por_clase_30m_1985_2026.csv"), encoding="utf-8-sig") as f:
        filas = list(csv.DictReader(f))
    nombres = [k for k in filas[0] if k != "year"]
    os.makedirs(os.path.join(SALIDA, "suelo"), exist_ok=True)
    caja, colores, anios = None, None, []
    for f in sorted(glob.glob(os.path.join(base, "clases_7", "clases7_*.tif"))):
        anio = int(os.path.basename(f)[8:12])
        with tifffile.TiffFile(f) as t:
            pg = t.pages[0]
            a = pg.asarray()
            cm = (pg.colormap // 257).astype(np.uint8)
            sx, sy = pg.tags["ModelPixelScaleTag"].value[:2]
            x0, y0 = pg.tags["ModelTiepointTag"].value[3:5]
        caja = [round(x0, 5), round(y0 - a.shape[0] * sy, 5), round(x0 + a.shape[1] * sx, 5), round(y0, 5)]
        colores = ["#%02x%02x%02x" % tuple(cm[:, i]) for i in range(len(nombres))]
        g = mayoria(a, 255, SUAVE)
        im = Image.fromarray(g, "P")
        im.putpalette([v for i in range(256) for v in (cm[:, i] if i < len(nombres) else (0, 0, 0))])
        im.save(os.path.join(SALIDA, "suelo", f"clases_{anio}.png"), optimize=True, transparency=255)
        anios.append(anio)
    tam = sum(os.path.getsize(p) for p in glob.glob(os.path.join(SALIDA, "suelo", "*.png")))
    print(f"  suelo: {len(anios)} mapas, {tam / 1e6:.1f} MB")
    return {"clases": [{"nombre": n, "color": c} for n, c in zip(nombres, colores)], "caja": caja, "mapas": anios,
            "km2": [[int(r["year"])] + [num(r[n], 1) for n in nombres] for r in filas]}


# ---------------------------------------------------------------------
# Camellones zenues y arqueologia (carpeta Camellones_Zenu_Geovisor, aparte del paquete)
# ---------------------------------------------------------------------
CAM = os.path.join(AQUI, "..", "Camellones_Zenu_Geovisor", "Camellones_Zenu_Geovisor")
CAPAS_CAM = [   # (shapefile, capa, campos que se conservan)
    ("01_Camellones/Camellones_zenues_poligonos", None, ["id_orig", "alt_cons_m", "area_m2", "fecha_img"]),
    ("01_Camellones/Camellones_ejes_IslaPerez_lineas", "eje", ["id_orig", "grupo", "long_m", "rumbo", "ang_cauce"]),
    ("01_Camellones/Cauces_antiguos_IslaPerez_lineas", "paleocauce", ["id_orig", "long_m"]),
    ("01_Camellones/Cauces_antiguos_SanJorge_poligonos", "paleocauce", []),
    ("02_Arqueologia/Plataformas_vivienda_SanJorge", "plataforma", ["area_m2"]),
    ("02_Arqueologia/Area_estudio_LaMojana_Vieri2026", "area_estudio", ["nombre", "area_ha"]),
    ("02_Arqueologia/Areas_prospectadas_MontejoRojas1992", "area_montejo", ["nombre", "precision"]),
    ("02_Arqueologia/Sitios_MontejoRojas1992_puntos", "sitio_montejo", ["sitio", "precision"]),
    ("02_Arqueologia/Zonas_revisadas_GoogleEarth_2026", "zona_revisada", ["zona", "resultado"]),
]


def camellones():
    a84 = Transformer.from_crs(9377, 4326, always_xy=True)
    feats, n = [], {}
    for rel, capa, campos in CAPAS_CAM:
        r = shapefile.Reader(os.path.join(CAM, "shapefiles_EPSG9377", rel + ".shp"), encoding="utf-8")
        for s, rec in zip(r.shapes(), r.records()):
            d = rec.as_dict()
            c = capa or ("camellon" if d["estado"].startswith("confirmado") else "franja")
            # a grados y redondeado a ~0,1 m (set_precision borraba algun camellon angosto)
            g = shapely.transform(shape(s.__geo_interface__).simplify(0.3),
                                  lambda xy: np.round(np.column_stack(a84.transform(xy[:, 0], xy[:, 1])), 6))
            p = {"capa": c}
            p.update({k: d[k] for k in campos if d.get(k) not in (None, "", "nan")})
            if c == "paleocauce":   # los de San Jorge (Vieri, poligonos) y los de Isla Perez (lineas a mano)
                p["region"] = "San Jorge" if "SanJorge" in rel else "Isla Pérez"
            feats.append((p, g))
            n[c] = n.get(c, 0) + 1
    guardar("camellones", feats)
    cam = [p for p, _ in feats if p["capa"] == "camellon"]
    alt = sorted(p["alt_cons_m"] for p in cam if p.get("alt_cons_m") is not None)
    ha = lambda capa: round(sum(p.get("area_m2", 0) for p, _ in feats if p["capa"] == capa) / 1e4, 1)
    # relieve LiDAR de 1 m (EPSG:3116) -> imagen en grados para ponerla encima del mapa
    with tifffile.TiffFile(os.path.join(CAM, "raster", "Relieve_local_LiDAR_camellones_SanMarcos_EPSG3116.tif")) as t:
        pg = t.pages[0]
        rgb = np.moveaxis(pg.asarray(), 0, -1)
        x0, y0 = pg.tags["ModelTiepointTag"].value[3:5]
        sx, sy = pg.tags["ModelPixelScaleTag"].value[:2]
    alto, ancho = rgb.shape[:2]
    a84b = Transformer.from_crs(3116, 4326, always_xy=True)
    lon, lat = a84b.transform([x0, x0 + ancho * sx, x0, x0 + ancho * sx], [y0, y0, y0 - alto * sy, y0 - alto * sy])
    caja = [min(lon), min(lat), max(lon), max(lat)]
    W = 1600
    H = int(W * (caja[3] - caja[1]) / (caja[2] - caja[0]))
    LO, LA = np.meshgrid(np.linspace(caja[0], caja[2], W), np.linspace(caja[3], caja[1], H))
    X, Y = Transformer.from_crs(4326, 3116, always_xy=True).transform(LO, LA)
    ci, fi = ((X - x0) / sx).astype(int), ((y0 - Y) / sy).astype(int)
    ok = (ci >= 0) & (ci < ancho) & (fi >= 0) & (fi < alto)
    sal = rgb[np.where(ok, fi, 0), np.where(ok, ci, 0)]
    sal[~ok] = 0
    os.makedirs(os.path.join(SALIDA, "camellones"), exist_ok=True)
    Image.fromarray(sal).save(os.path.join(SALIDA, "camellones", "lidar_sanmarcos.jpg"), quality=82)
    print(f"  camellones: {n}; LiDAR {W}x{H}")
    return {"n": n, "ha_camellones": ha("camellon"), "ha_franjas": ha("franja"), "ha_plataformas": ha("plataforma"),
            "alt_mediana": round(alt[len(alt) // 2], 2), "alt": [round(v, 3) for v in alt],
            "lidar": {"url": "datos/camellones/lidar_sanmarcos.jpg", "caja": [round(v, 6) for v in caja]}}


# ---------------------------------------------------------------------
# Indice espectral de turbidez (NDTI) y color real en la desembocadura, 1985-2026
# ---------------------------------------------------------------------
NDTI_RANGO = (-0.4, 0.2)    # rango sugerido en LEEME_desembocadura.txt
NDTI_PASOS = 24             # la imagen se guarda con paleta: mucho mas liviana
NDTI_PALETA = [(255, 247, 188), (254, 196, 79), (217, 95, 14), (102, 37, 6)]   # amarillo claro a cafe
SENSOR_S2 = 2015            # desde 2015 es Sentinel-2: el nivel del NDTI cambia (no comparable con Landsat)
RED = 2                     # de 20 m a 40 m (Landsat es de 30 m)


def rampa(t):
    t = np.clip(t, 0, 1) * (len(NDTI_PALETA) - 1)
    i = np.minimum(t.astype(int), len(NDTI_PALETA) - 2)
    f = (t - i)[..., None]
    a, b = np.array(NDTI_PALETA)[i], np.array(NDTI_PALETA)[i + 1]
    return (a + (b - a) * f).astype(np.uint8)


def desembocadura():
    base = os.path.join(PAQ, "11_Desembocadura_sedimentos")
    os.makedirs(os.path.join(SALIDA, "desembocadura"), exist_ok=True)
    lo, hi = NDTI_RANGO
    paleta = rampa((np.arange(NDTI_PASOS) + 0.5) / NDTI_PASOS)
    anios, caja = [], None
    for f in sorted(glob.glob(os.path.join(base, "NDTI", "NDTI_*_desembocadura_sinu.tif"))):
        anio = int(os.path.basename(f)[5:9])
        with tifffile.TiffFile(f) as t:
            pg = t.pages[0]
            v = pg.asarray().astype(np.float32)
            sx, sy = pg.tags["ModelPixelScaleTag"].value[:2]
            x0, y0 = pg.tags["ModelTiepointTag"].value[3:5]
        caja = [round(x0, 5), round(y0 - v.shape[0] * sy, 5), round(x0 + v.shape[1] * sx, 5), round(y0, 5)]
        ok = v > -9000
        agua = v[ok]
        # imagen: indice de paleta por pixel, 255 = sin dato (transparente)
        h, w = (v.shape[0] // RED) * RED, (v.shape[1] // RED) * RED
        vr = v[:h:RED, :w:RED]
        idx = np.clip(((vr - lo) / (hi - lo) * NDTI_PASOS).astype(int), 0, NDTI_PASOS - 1).astype(np.uint8)
        idx[vr <= -9000] = 255
        im = Image.fromarray(idx, "P")
        im.putpalette(paleta.flatten().tolist() + [0, 0, 0] * (256 - NDTI_PASOS))
        im.save(os.path.join(SALIDA, "desembocadura", f"ndti_{anio}.png"), optimize=True, transparency=255)
        # color real: reflectancia x 10000 estirada entre 0 y 1500 (el agua es oscura)
        with tifffile.TiffFile(os.path.join(base, "RGB", f"RGB_{anio}_desembocadura_sinu.tif")) as t:
            rgb = t.pages[0].asarray().astype(np.float32)
        if rgb.shape[0] == 3:
            rgb = np.moveaxis(rgb, 0, -1)
        sin = (rgb <= -9000).any(axis=-1)
        rgb = (255 * np.clip(rgb / 1500, 0, 1) ** 0.8).astype(np.uint8)
        rgb[sin] = 0
        Image.fromarray(rgb[:h:RED, :w:RED]).save(os.path.join(SALIDA, "desembocadura", f"rgb_{anio}.jpg"), quality=80)
        anios.append({"anio": anio, "mediana": num(float(np.median(agua)), 3) if agua.size else None,
                      "p90": num(float(np.percentile(agua, 90)), 3) if agua.size else None,
                      "cobertura": num(100 * ok.mean(), 1), "rgb": num(100 * (~sin).mean(), 1),
                      "sensor": "Sentinel-2" if anio >= SENSOR_S2 else "Landsat"})
    tam = sum(os.path.getsize(p) for p in glob.glob(os.path.join(SALIDA, "desembocadura", "*")))
    print(f"  desembocadura: {len(anios)} anos, {tam / 1e6:.1f} MB")
    return {"caja": caja, "rango": list(NDTI_RANGO), "paleta": ["#%02x%02x%02x" % tuple(c) for c in NDTI_PALETA],
            "sensor_s2": SENSOR_S2, "anios": anios}


# ---------------------------------------------------------------------
# Mapas historicos: carta de Bolivar 1864 y planchas INCORA 1963 (bloque 5 del StoryMap)
# ---------------------------------------------------------------------
HIS = os.path.join(AQUI, "..", "Mapas_Historicos_Agua_como_archivo", "Mapas_Historicos_Agua_como_archivo")
MAPAS_HIS = [   # (clave, GeoTIFF en EPSG:9377, recorte en grados o None, metros por pixel de salida, formato)
    ("bolivar_1864", "Bolivar_1864_9377.tif", (-76.3, 8.3, -75.25, 9.5), 70, "jpg"),
    ("incora_61-IB-4d", "INCORA1963_61-IB-4d_9377.tif", None, 2, "webp"),
    ("incora_61-IB-2b", "INCORA1963_61-IB-2b_9377.tif", None, 2, "webp"),
]


def historicos():
    """Pasa los GeoTIFF a imagenes en grados (para image sources de MapLibre) y las cienagas de 1963 a geojson."""
    os.makedirs(os.path.join(SALIDA, "historicos"), exist_ok=True)
    a84, a9377 = Transformer.from_crs(9377, 4326, always_xy=True), Transformer.from_crs(4326, 9377, always_xy=True)
    mapas = {}
    for clave, tif, recorte, paso, fmt in MAPAS_HIS:
        with tifffile.TiffFile(os.path.join(HIS, "02_Georreferenciados_EPSG9377", tif)) as t:
            pg = t.pages[0]
            rgb = pg.asarray()
            x0, y0 = pg.tags["ModelTiepointTag"].value[3:5]
            s = pg.tags["ModelPixelScaleTag"].value[0]
        alto, ancho = rgb.shape[:2]
        if recorte is None:
            lon, lat = a84.transform([x0, x0 + ancho * s, x0, x0 + ancho * s], [y0, y0, y0 - alto * s, y0 - alto * s])
            recorte = (min(lon), min(lat), max(lon), max(lat))
        caja = [round(v, 6) for v in recorte]
        # tamano de salida segun los metros por pixel pedidos (1 grado ~ 111 km)
        W = int((caja[2] - caja[0]) * 111320 * np.cos(np.radians((caja[1] + caja[3]) / 2)) / paso)
        H = int((caja[3] - caja[1]) * 110570 / paso)
        LO, LA = np.meshgrid(np.linspace(caja[0], caja[2], W), np.linspace(caja[3], caja[1], H))
        X, Y = a9377.transform(LO, LA)
        ci, fi = ((X - x0) / s).astype(int), ((y0 - Y) / s).astype(int)
        ok = (ci >= 0) & (ci < ancho) & (fi >= 0) & (fi < alto)
        sal = rgb[np.where(ok, fi, 0), np.where(ok, ci, 0)]
        # las planchas traen las esquinas giradas en negro: se vuelven transparentes
        ok &= sal.astype(int).sum(axis=2) > 0
        ruta = os.path.join(SALIDA, "historicos", f"{clave}.{fmt}")
        if fmt == "webp":
            Image.fromarray(np.dstack([sal, np.where(ok, 255, 0).astype(np.uint8)])).save(ruta, quality=72, method=6)
        else:
            Image.fromarray(sal).save(ruta, quality=80)
        mapas[clave] = {"url": f"datos/historicos/{clave}.{fmt}", "caja": caja}
        print(f"  historicos/{clave}: {W}x{H}, {os.path.getsize(ruta) / 1e6:.2f} MB")
    feats = []
    for plancha in ("61-IB-4d", "61-IB-2b"):
        r = shapefile.Reader(os.path.join(HIS, "03_Cienagas_1963_EPSG9377", f"Cienagas_1963_{plancha}"), encoding="utf-8")
        for sh in r.shapes():
            g = shape(sh.__geo_interface__)
            ha = g.area / 1e4
            g = shapely.transform(g.simplify(0.5), lambda xy: np.round(np.column_stack(a84.transform(xy[:, 0], xy[:, 1])), 6))
            feats.append(({"plancha": plancha, "ha": round(ha, 2)}, g))
    guardar("cienagas_1963", feats)
    ha = {p: round(sum(f[0]["ha"] for f in feats if f[0]["plancha"] == p), 1) for p in ("61-IB-4d", "61-IB-2b")}
    return {"mapas": mapas, "cienagas": {"n": len(feats), "ha": ha}}


# ---------------------------------------------------------------------
# Indices espectrales NDVI, NDWI y MNDWI 1985-2026 (cuenca, zonas IDEAM, 58 subcuencas y rasteres)
# ---------------------------------------------------------------------
IDX = os.path.join(AQUI, "..", "Indices_Espectrales_Geovisor")
IDX_NOMBRES = ["NDVI", "NDWI", "MNDWI"]
IDX_RANGO = {"NDVI": (0.2, 0.9), "NDWI": (-0.7, 0.3), "MNDWI": (-0.7, 0.3)}   # los del LEEME
IDX_PALETA = {"NDVI": [(140, 81, 10), (223, 194, 125), (166, 217, 106), (26, 150, 65), (0, 68, 27)],      # cafe a verde
              "NDWI": [(140, 81, 10), (216, 179, 101), (246, 232, 195), (146, 197, 222), (33, 102, 172), (5, 48, 97)]}   # cafe a azul
IDX_PALETA["MNDWI"] = IDX_PALETA["NDWI"]
IDX_PASOS = 24
IDX_PASO_M = 250            # los rasteres vienen a 90 m: se guardan a 250 m (unos 90 KB por imagen, ~12 MB las 126)
IDX_SENSOR = 2013           # de 2012 a 2013 cambia el sensor (TM/ETM+ a OLI): el NDVI salta aunque se armonice


def sin_tildes(t):
    import unicodedata
    return "".join(c for c in unicodedata.normalize("NFD", t or "") if unicodedata.category(c) != "Mn").lower().strip()


def indices():
    """Tablas de la carpeta Indices_Espectrales_Geovisor a series.json y rasteres anuales a PNG con paleta en grados."""
    def leer_csv(nombre):
        with open(os.path.join(IDX, nombre), encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))
    cuenca = leer_csv("indices_cuenca_anual_1985-2026.csv")
    campos = ["NDVI", "NDWI", "MNDWI", "EVI", "NDMI", "NBR", "NDBI", "BSI"]
    anios = [int(r["anio"]) for r in cuenca]
    salida = {"anios": anios, "sensor": IDX_SENSOR, "campos": campos,
              "cuenca": [[num(r[c], 4) for c in campos] for r in cuenca],
              "escenas": [int(r["n_escenas"]) for r in cuenca], "ventana": [int(r["ventana_meses"]) for r in cuenca],
              "sensores": [r["sensor"] for r in cuenca], "notas": {r["anio"]: r["nota"] for r in cuenca if r["nota"]}}

    def por_subcuenca(nombre):
        filas = {}
        for r in leer_csv(nombre):
            filas.setdefault(r["subcuenca"], {"km2": num(r["km2"], 1), "v": {}})["v"][int(r["year"])] = \
                [num(r[c], 4) for c in IDX_NOMBRES] + [num(r["cobertura"], 3)]
        return {k: {"km2": d["km2"], "v": [d["v"].get(a, [None] * 4) for a in anios]} for k, d in filas.items()}
    salida["zonas"] = por_subcuenca("indices_subcuencas_IDEAM_1985-2026.csv")
    # las 58 subcuencas del panel se unen por nombre con las del motor 3D (datos/subcuencas.geojson)
    sub = json.load(open(os.path.join(SALIDA, "subcuencas.geojson"), encoding="utf-8"))
    ids = {sin_tildes(f["properties"]["nombre"]): f["properties"]["id"] for f in sub["features"]}
    tabla = por_subcuenca("indices_subcuencas_1985-2026.csv")
    salida["sub"] = {ids[sin_tildes(k)]: d["v"] for k, d in tabla.items() if sin_tildes(k) in ids}
    faltan = [k for k in tabla if sin_tildes(k) not in ids]
    print(f"  indices: {len(salida['sub'])} subcuencas unidas" + (f", sin unir: {faltan}" if faltan else ""))

    # rasteres: EPSG:32618 a grados por muestreo inverso, un PNG con paleta por indice y ano (255 = sin dato)
    os.makedirs(os.path.join(SALIDA, "indices"), exist_ok=True)
    a84, a32618 = Transformer.from_crs(32618, 4326, always_xy=True), Transformer.from_crs(4326, 32618, always_xy=True)
    rasteres, caja, malla = [], None, None
    for f in sorted(glob.glob(os.path.join(IDX, "rasteres_anuales", "Sinu_NDVI_NDWI_MNDWI_*.tif"))):
        anio = int(os.path.basename(f)[-8:-4])
        with tifffile.TiffFile(f) as t:
            pg = t.pages[0]
            v = pg.asarray()
            x0, y0 = pg.tags["ModelTiepointTag"].value[3:5]
            s = pg.tags["ModelPixelScaleTag"].value[0]
        alto, ancho = v.shape[:2]
        if malla is None:
            lon, lat = a84.transform([x0, x0 + ancho * s, x0, x0 + ancho * s], [y0, y0, y0 - alto * s, y0 - alto * s])
            caja = [round(min(lon), 5), round(min(lat), 5), round(max(lon), 5), round(max(lat), 5)]
            W = int((caja[2] - caja[0]) * 111320 * np.cos(np.radians((caja[1] + caja[3]) / 2)) / IDX_PASO_M)
            H = int((caja[3] - caja[1]) * 110570 / IDX_PASO_M)
            LO, LA = np.meshgrid(np.linspace(caja[0], caja[2], W), np.linspace(caja[3], caja[1], H))
            X, Y = a32618.transform(LO, LA)
            ci, fi = ((X - x0) / s).astype(int), ((y0 - Y) / s).astype(int)
            ok = (ci >= 0) & (ci < ancho) & (fi >= 0) & (fi < alto)
            malla = (np.where(ok, fi, 0), np.where(ok, ci, 0), ok)
        fi, ci, ok = malla
        for b, nombre in enumerate(IDX_NOMBRES):
            m = v[fi, ci, b]
            lo, hi = IDX_RANGO[nombre]
            idx = np.clip(((m / 1e4 - lo) / (hi - lo) * IDX_PASOS).astype(int), 0, IDX_PASOS - 1).astype(np.uint8)
            idx[(m == -32768) | ~ok] = 255
            im = Image.fromarray(idx, "P")
            pal = IDX_PALETA[nombre]
            t_ = np.clip((np.arange(IDX_PASOS) + 0.5) / IDX_PASOS, 0, 1) * (len(pal) - 1)
            i_ = np.minimum(t_.astype(int), len(pal) - 2)
            col = (np.array(pal)[i_] + (np.array(pal)[i_ + 1] - np.array(pal)[i_]) * (t_ - i_)[:, None]).astype(np.uint8)
            im.putpalette(col.flatten().tolist() + [0, 0, 0] * (256 - IDX_PASOS))
            im.save(os.path.join(SALIDA, "indices", f"{nombre.lower()}_{anio}.png"), optimize=True, transparency=255)
        rasteres.append(anio)
    tam = sum(os.path.getsize(p) for p in glob.glob(os.path.join(SALIDA, "indices", "*")))
    print(f"  indices: {len(rasteres)} anos x 3 imagenes ({W}x{H}), {tam / 1e6:.1f} MB")
    salida["raster"] = {"caja": caja, "anios": rasteres, "rango": {k: list(v) for k, v in IDX_RANGO.items()},
                        "paleta": {k: ["#%02x%02x%02x" % c for c in p] for k, p in IDX_PALETA.items()}}
    return salida


# ---------------------------------------------------------------------
def main():
    solo = {"suelo": suelo, "camellones": camellones, "desembocadura": desembocadura, "historicos": historicos,
            "indices": indices}
    if set(sys.argv[1:]) & set(solo):   # solo esas partes, sin rehacer el resto
        ruta = os.path.join(SALIDA, "series.json")
        todo = json.load(open(ruta, encoding="utf-8"))
        for k in sys.argv[1:]:
            todo[k] = solo[k]()
        open(ruta, "w", encoding="utf-8").write(json.dumps(todo, ensure_ascii=False, separators=(",", ":")))
        return
    cuenca_reg = leer("01_Cuenca_Drenaje_DEM/Cuenca_completa/Cuenca_Sinu.shp")
    cuenca = unary_union([g for _, g in cuenca_reg])
    shp = os.path.join(PAQ, "01_Cuenca_Drenaje_DEM/Cuenca_completa/Cuenca_Sinu.shp")
    cuenca_9377 = unary_union([shape(s.__geo_interface__) for s in shapefile.Reader(shp).shapes()])

    print("Capas tematicas")
    series = capas(cuenca, round(cuenca_9377.area / 1e6, 1))
    print("Inundacion y sequia")
    info_agua = agua(cuenca_9377)

    def csv_rel(rel):
        with open(os.path.join(PAQ, rel), encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))
    area = [[int(r["year"]), num(r["area_inundacion_max_km2"], 2), num(r["area_sequia_km2"], 2)]
            for r in csv_rel("05_Grafica_tiempo_area_agua/area_inundacion_sequia_anual_1984_2026.csv")]
    oni = [[int(r["year"]), int(r["month"]), num(r["oni"]), num(r["roni"])] for r in csv_rel("08_ONI_RONI/oni_roni_mensual.csv")]
    lot = [{"id": r["sub_id"], "nombre": arreglar(r["nombre"]), "km2": num(r["km2"], 1), "ha_lotico": num(r["ha_lotico"], 1),
            "ha_lentico": num(r["ha_lentico"], 1), "ha_agua": num(r["ha_agua"], 1), "pct_agua": num(r["pct_agua_subcuenca"], 2),
            "ha_permanente": num(r["ha_permanente"], 1), "ha_estacional": num(r["ha_estacional"], 1),
            "n_lotico": int(float(r["n_lotico"] or 0)), "n_lentico": int(float(r["n_lentico"] or 0))}
           for r in csv_rel("07_Lotico_Lentico/2020/resultados_por_subcuenca.csv")]
    print("Cobertura del suelo")
    todo = {"estaciones": series, "area_agua": area, "oni_roni": oni, "lotico_lentico": lot, "agua": info_agua, "suelo": suelo(),
            "camellones": camellones(), "desembocadura": desembocadura(), "historicos": historicos(),
            "indices": indices()}
    texto = json.dumps(todo, ensure_ascii=False, separators=(",", ":"))
    open(os.path.join(SALIDA, "series.json"), "w", encoding="utf-8").write(texto)
    print(f"  series.json: {len(texto) / 1e6:.2f} MB")

    print("Motor 3D")
    motor3d(cuenca, info_agua)


if __name__ == "__main__":
    main()
